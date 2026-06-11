#!/usr/bin/env python3
"""Generate an International Foreign Travel Brief (Thailand) as HTML, PDF, and XLSX.

Used by the FSO for cleared and non-cleared foreign travel reporting.
"""

import os

OUT_DIR = os.path.dirname(os.path.abspath(__file__))

# ---------------------------------------------------------------- brief data
COUNTRY = "Thailand"
DEPART = "6/18/2026"
RETURN = "7/6/2026"
DATE_ISSUED = "6/11/2026"

LINKS = [
    ("Register with STEP (Smart Traveler Enrollment Program)", "https://step.state.gov/"),
    ("Country Information", "https://travel.state.gov/content/travel/en/international-travel/International-Travel-Country-Information-Pages/Thailand.html"),
    ("Country Security Report (OSAC)", "https://www.osac.gov/Country/Thailand/Detail"),
    ("Fact Sheet", "https://www.state.gov/countries-areas/thailand/"),
    ("Assistance for U.S. Citizens in Thailand", "https://th.usembassy.gov/"),
]

EMBASSY = [
    ("Embassy", "U.S. Embassy Bangkok"),
    ("Address", "95 Wireless Road, Bangkok 10330, Thailand"),
    ("Telephone", "+ (66) (2) 205-4049"),
    ("Emergency", "+ (66) (2) 205-4000"),
    ("Fax", "+ (66) (2) 205-4103"),
    ("Email", "ACSBkk@state.gov"),
    ("Website", "https://th.usembassy.gov/embassy-consulate/bangkok/"),
]

GREETING = "Hi,"
INTRO = (f"Thank you for reporting your upcoming travel to {COUNTRY} and return "
         f"from {DEPART} through {RETURN}.")
REVIEW = ("Please, review and familiarize yourself with the updated resources linked "
          "below prior to your trip. These links contain important information regarding "
          "the current status and conditions in the country you will be visiting:")
REMINDER = "Reminder: you must complete the post-travel debrief within five (5) days of return."
CLOSING = "Please confirm receipt and safe travels."
TITLE = "International Foreign Travel Brief"
SUBTITLE = "Cleared & Non-Cleared Foreign Travel Reporting — Facility Security Officer (FSO)"

BASENAME = os.path.join(OUT_DIR, "Foreign_Travel_Brief_Thailand")

# ------------------------------------------------------------------ HTML ---
html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{TITLE} — {COUNTRY}</title>
<style>
  body {{ font-family: "Segoe UI", Arial, sans-serif; color: #1a1a2e; margin: 0;
         background: #f4f6f8; }}
  .page {{ max-width: 820px; margin: 24px auto; background: #fff; padding: 36px 48px;
           box-shadow: 0 2px 10px rgba(0,0,0,.12); border-top: 6px solid #1f3a93; }}
  header h1 {{ color: #1f3a93; margin: 0 0 4px; font-size: 26px; }}
  header p.sub {{ margin: 0; color: #555; font-size: 13px; font-style: italic; }}
  .meta {{ margin: 18px 0; border-collapse: collapse; width: 100%; }}
  .meta th, .meta td {{ border: 1px solid #c9d2e3; padding: 7px 10px; font-size: 13px;
        text-align: left; }}
  .meta th {{ background: #1f3a93; color: #fff; width: 200px; }}
  h2 {{ color: #1f3a93; font-size: 17px; border-bottom: 2px solid #1f3a93;
        padding-bottom: 4px; margin-top: 28px; }}
  ul.links li {{ margin: 8px 0; font-size: 14px; }}
  a {{ color: #1158a7; }}
  .embassy td {{ border: 1px solid #c9d2e3; padding: 7px 10px; font-size: 13px; }}
  .embassy td:first-child {{ background: #eef2fa; font-weight: 600; width: 200px; }}
  .embassy {{ border-collapse: collapse; width: 100%; }}
  .reminder {{ background: #fff3cd; border: 1px solid #ffc107; border-left: 6px solid #ffc107;
        padding: 12px 16px; margin: 24px 0; font-weight: 600; font-size: 14px; }}
  .sig {{ margin-top: 32px; font-size: 14px; }}
  footer {{ margin-top: 36px; font-size: 11px; color: #888; border-top: 1px solid #ddd;
        padding-top: 10px; text-align: center; }}
</style>
</head>
<body>
<div class="page">
  <header>
    <h1>{TITLE}</h1>
    <p class="sub">{SUBTITLE}</p>
  </header>

  <table class="meta">
    <tr><th>Destination Country</th><td>{COUNTRY}</td></tr>
    <tr><th>Departure Date</th><td>{DEPART}</td></tr>
    <tr><th>Return Date</th><td>{RETURN}</td></tr>
    <tr><th>Brief Issued</th><td>{DATE_ISSUED}</td></tr>
    <tr><th>Post-Travel Debrief Due</th><td>Within five (5) days of return</td></tr>
  </table>

  <p>{GREETING}</p>
  <p>{INTRO}</p>
  <p>{REVIEW}</p>

  <h2>Required Pre-Travel Resources</h2>
  <ul class="links">
""" + "\n".join(
    f'    <li><strong>{label}:</strong> <a href="{url}">{url}</a></li>'
    for label, url in LINKS
) + f"""
  </ul>

  <h2>Assistance for U.S. Citizens</h2>
  <table class="embassy">
""" + "\n".join(
    '    <tr><td>{}</td><td>{}</td></tr>'.format(
        k,
        f'<a href="{v}">{v}</a>' if v.startswith("http")
        else (f'<a href="mailto:{v}">{v}</a>' if "@" in v else v))
    for k, v in EMBASSY
) + f"""
  </table>

  <div class="reminder">&#9888;&nbsp; {REMINDER}</div>

  <p>{CLOSING}</p>
  <p class="sig">Best,<br><br>____________________________<br>Facility Security Officer (FSO)</p>

  <footer>International Foreign Travel Brief &mdash; {COUNTRY} &mdash; Travel {DEPART} to {RETURN}</footer>
</div>
</body>
</html>
"""

with open(BASENAME + ".html", "w", encoding="utf-8") as f:
    f.write(html)
print("HTML written")

# ------------------------------------------------------------------ XLSX ---
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Travel Brief"

navy = "1F3A93"
hdr_fill = PatternFill("solid", fgColor=navy)
lbl_fill = PatternFill("solid", fgColor="EEF2FA")
warn_fill = PatternFill("solid", fgColor="FFF3CD")
thin = Side(style="thin", color="C9D2E3")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

ws.column_dimensions["A"].width = 32
ws.column_dimensions["B"].width = 95

def row(r, a, b="", a_bold=False, fill=None, merge=False):
    ws.cell(row=r, column=1, value=a)
    if merge:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    else:
        ws.cell(row=r, column=2, value=b)
    for c in (1, 2):
        cell = ws.cell(row=r, column=c)
        cell.alignment = wrap
        cell.border = border
        if fill:
            cell.fill = fill
        if a_bold and c == 1:
            cell.font = Font(bold=True)

r = 1
ws.merge_cells("A1:B1")
ws["A1"] = TITLE
ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
ws["A1"].fill = hdr_fill
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 28

ws.merge_cells("A2:B2")
ws["A2"] = SUBTITLE
ws["A2"].font = Font(italic=True, size=10)
ws["A2"].alignment = Alignment(horizontal="center")

r = 4
for label, val in [("Destination Country", COUNTRY), ("Departure Date", DEPART),
                   ("Return Date", RETURN), ("Brief Issued", DATE_ISSUED),
                   ("Post-Travel Debrief Due", "Within five (5) days of return")]:
    row(r, label, val, a_bold=True, fill=lbl_fill)
    ws.cell(row=r, column=2).fill = PatternFill()
    ws.cell(row=r, column=2).border = border
    r += 1

r += 1
row(r, GREETING, merge=True); r += 1
row(r, INTRO, merge=True); ws.row_dimensions[r].height = 30; r += 1
row(r, REVIEW, merge=True); ws.row_dimensions[r].height = 45; r += 1

r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
c = ws.cell(row=r, column=1, value="REQUIRED PRE-TRAVEL RESOURCES")
c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill
ws.cell(row=r, column=2).fill = hdr_fill
r += 1
for label, url in LINKS:
    ws.cell(row=r, column=1, value=label).font = Font(bold=True)
    link = ws.cell(row=r, column=2, value=url)
    link.hyperlink = url
    link.font = Font(color="1158A7", underline="single")
    for col in (1, 2):
        ws.cell(row=r, column=col).border = border
        ws.cell(row=r, column=col).alignment = wrap
    ws.cell(row=r, column=1).fill = lbl_fill
    r += 1

r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
c = ws.cell(row=r, column=1, value="ASSISTANCE FOR U.S. CITIZENS")
c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill
ws.cell(row=r, column=2).fill = hdr_fill
r += 1
for k, v in EMBASSY:
    ws.cell(row=r, column=1, value=k).font = Font(bold=True)
    cell = ws.cell(row=r, column=2, value=v)
    if v.startswith("http"):
        cell.hyperlink = v
        cell.font = Font(color="1158A7", underline="single")
    elif "@" in v:
        cell.hyperlink = "mailto:" + v
        cell.font = Font(color="1158A7", underline="single")
    for col in (1, 2):
        ws.cell(row=r, column=col).border = border
        ws.cell(row=r, column=col).alignment = wrap
    ws.cell(row=r, column=1).fill = lbl_fill
    r += 1

r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
c = ws.cell(row=r, column=1, value="⚠ " + REMINDER)
c.font = Font(bold=True); c.fill = warn_fill
ws.cell(row=r, column=2).fill = warn_fill
for col in (1, 2):
    ws.cell(row=r, column=col).border = border
ws.row_dimensions[r].height = 24
r += 2

row(r, CLOSING, merge=True); r += 2
row(r, "Best,", merge=True); r += 2
row(r, "____________________________", merge=True); r += 1
row(r, "Facility Security Officer (FSO)", merge=True)

wb.save(BASENAME + ".xlsx")
print("XLSX written")

# ------------------------------------------------------------------- PDF ---
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, HRFlowable)

NAVY = colors.HexColor("#1F3A93")
LIGHT = colors.HexColor("#EEF2FA")
GRID = colors.HexColor("#C9D2E3")
WARN_BG = colors.HexColor("#FFF3CD")
WARN_BD = colors.HexColor("#FFC107")

styles = getSampleStyleSheet()
title_st = ParagraphStyle("T", parent=styles["Title"], textColor=NAVY, fontSize=20,
                          spaceAfter=2)
sub_st = ParagraphStyle("S", parent=styles["Normal"], fontSize=9.5,
                        textColor=colors.HexColor("#555555"), alignment=1,
                        fontName="Helvetica-Oblique", spaceAfter=14)
h2_st = ParagraphStyle("H2", parent=styles["Heading2"], textColor=NAVY, fontSize=12.5,
                       spaceBefore=14, spaceAfter=6)
body_st = ParagraphStyle("B", parent=styles["Normal"], fontSize=10.5, leading=15,
                         spaceAfter=8)
cell_st = ParagraphStyle("C", parent=styles["Normal"], fontSize=9.5, leading=13)
cellb_st = ParagraphStyle("CB", parent=cell_st, fontName="Helvetica-Bold")
warn_st = ParagraphStyle("W", parent=body_st, fontName="Helvetica-Bold", fontSize=10.5)

doc = SimpleDocTemplate(BASENAME + ".pdf", pagesize=letter,
                        leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                        topMargin=0.7 * inch, bottomMargin=0.7 * inch,
                        title=f"{TITLE} - {COUNTRY}")
el = []
el.append(Paragraph(TITLE, title_st))
el.append(Paragraph(SUBTITLE, sub_st))
el.append(HRFlowable(width="100%", thickness=2.5, color=NAVY, spaceAfter=12))

meta = [[Paragraph(k, cellb_st), Paragraph(v, cell_st)] for k, v in [
    ("Destination Country", COUNTRY), ("Departure Date", DEPART),
    ("Return Date", RETURN), ("Brief Issued", DATE_ISSUED),
    ("Post-Travel Debrief Due", "Within five (5) days of return")]]
t = Table(meta, colWidths=[1.9 * inch, 4.9 * inch])
t.setStyle(TableStyle([
    ("GRID", (0, 0), (-1, -1), 0.6, GRID),
    ("BACKGROUND", (0, 0), (0, -1), LIGHT),
    ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
    ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
]))
el.append(t)
el.append(Spacer(1, 14))

el.append(Paragraph(GREETING, body_st))
el.append(Paragraph(INTRO, body_st))
el.append(Paragraph(REVIEW, body_st))

el.append(Paragraph("Required Pre-Travel Resources", h2_st))
links_rows = [[Paragraph(label, cellb_st),
               Paragraph(f'<link href="{url}" color="#1158A7"><u>{url}</u></link>', cell_st)]
              for label, url in LINKS]
t = Table(links_rows, colWidths=[2.2 * inch, 4.6 * inch])
t.setStyle(TableStyle([
    ("GRID", (0, 0), (-1, -1), 0.6, GRID),
    ("BACKGROUND", (0, 0), (0, -1), LIGHT),
    ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
    ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
]))
el.append(t)

el.append(Paragraph("Assistance for U.S. Citizens", h2_st))
emb_rows = []
for k, v in EMBASSY:
    if v.startswith("http"):
        val = Paragraph(f'<link href="{v}" color="#1158A7"><u>{v}</u></link>', cell_st)
    elif "@" in v:
        val = Paragraph(f'<link href="mailto:{v}" color="#1158A7"><u>{v}</u></link>', cell_st)
    else:
        val = Paragraph(v, cell_st)
    emb_rows.append([Paragraph(k, cellb_st), val])
t = Table(emb_rows, colWidths=[1.9 * inch, 4.9 * inch])
t.setStyle(TableStyle([
    ("GRID", (0, 0), (-1, -1), 0.6, GRID),
    ("BACKGROUND", (0, 0), (0, -1), LIGHT),
    ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
    ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
]))
el.append(t)
el.append(Spacer(1, 14))

warn = Table([[Paragraph(REMINDER, warn_st)]], colWidths=[6.8 * inch])
warn.setStyle(TableStyle([
    ("BACKGROUND", (0, 0), (-1, -1), WARN_BG),
    ("BOX", (0, 0), (-1, -1), 1, WARN_BD),
    ("LINEBEFORE", (0, 0), (0, -1), 4, WARN_BD),
    ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 12),
    ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
]))
el.append(warn)
el.append(Spacer(1, 14))

el.append(Paragraph(CLOSING, body_st))
el.append(Spacer(1, 10))
el.append(Paragraph("Best,", body_st))
el.append(Spacer(1, 24))
el.append(Paragraph("____________________________", body_st))
el.append(Paragraph("Facility Security Officer (FSO)", body_st))

doc.build(el)
print("PDF written")
