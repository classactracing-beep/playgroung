"""Document builders: SPP, policies, procedure packages, matrix, binder."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from . import branding as br
from .practices import DOMAINS, practices, practices_by_domain
from .procedures_data import (DOMAIN_GOVERNANCE, DOMAIN_GUIDANCE, PACKAGES,
                              POLICY_NAMES, evidence_for)
from .references import COMPLIANCE_REFERENCES, COMPLIANCE_STATEMENT
from .profile import fill
from . import spp_content

SAFE = str.maketrans({" ": "_", "/": "-", "&": "and", ",": ""})


def _safe(name):
    return name.translate(SAFE)


def _ensure(path):
    os.makedirs(path, exist_ok=True)
    return path


def _blocks(doc, blocks, profile):
    for kind, content in blocks:
        if kind == "p":
            doc.add_paragraph(fill(content, profile), style="Body Text"
                              if "Body Text" in [s.name for s in doc.styles]
                              else None)
        elif kind == "ul":
            for item in content:
                doc.add_paragraph(fill(item, profile), style="List Bullet")
        elif kind == "h2":
            doc.add_heading(fill(content, profile), level=2)
        elif kind == "sig":
            br.add_signature_block(
                doc, [(fill(n, profile), fill(t, profile)) for n, t in content])


# ---------------------------------------------------------------- SPP

def build_spp(profile, out_dir, include_appendices=True):
    """Branded SPP modeled on the AVASO Federal Solutions SPP 2026."""
    doc = br.new_document()
    title = "STANDARD PRACTICES AND PROCEDURES (SPP)"
    br.add_cover_page(doc, profile, title,
                      subtitle=f"CMMC Level 2 Security Program | "
                               f"{profile['document_year']}")
    br.add_disclosure_statement(doc, profile)
    br.add_page_break(doc)

    # Forward
    h = doc.add_paragraph()
    r = h.add_run("Forward")
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = br.HEADING_COLOR
    _blocks(doc, spp_content.FORWARD, profile)
    br.add_page_break(doc)

    br.add_toc(doc)
    br.add_page_break(doc)

    for heading, blocks in spp_content.SECTIONS:
        doc.add_heading(fill(heading, profile), level=1)
        _blocks(doc, blocks, profile)

    if include_appendices:
        br.add_page_break(doc)
        _appendix_a(doc, profile)
        br.add_page_break(doc)
        _appendix_b(doc, profile)

    br.add_header(doc, profile, "Standard Practices and Procedures")
    br.add_page_number_footer(
        doc, f"{profile['company_short_name']} SPP | Version "
             f"{profile['spp_version']} | Company Sensitive")
    _numbered_headings(doc)

    path = os.path.join(_ensure(out_dir),
                        f"{_safe(profile['company_short_name'])}_SPP_"
                        f"{profile['document_year']}.docx")
    doc.save(path)
    return path


def _numbered_headings(doc):
    """Apply sequential section numbers to Heading 1 paragraphs (skip appendices)."""
    n = 0
    for p in doc.paragraphs:
        if p.style.name == "Heading 1" and p.text and not p.text.startswith("Appendix"):
            n += 1
            for run in p.runs:
                run.text = run.text  # keep runs intact
            p.runs[0].text = f"{n}. {p.runs[0].text}" if p.runs else p.text
            if not p.runs:
                p.text = f"{n}. {p.text}"


def _matrix_rows(profile):
    rows = []
    for p in practices():
        role, freq = DOMAIN_GOVERNANCE[p["domain"]]
        rows.append([
            p["domain_name"],
            p["cmmc_id"],
            p["nist_id"],
            p["name"],
            POLICY_NAMES[p["domain"]],
            PACKAGES[p["domain"]]["procedure"]["title"],
            evidence_for(p["domain"]),
            role,
            freq,
        ])
    return rows


MATRIX_HEADERS = ["CMMC Domain", "CMMC Practice ID", "NIST SP 800-171 Rev. 2 ID",
                  "Requirement Title", "Policy Document", "Procedure Document",
                  "Evidence Artifact", "Responsible Role", "Review Frequency"]


def _appendix_a(doc, profile):
    doc.add_heading("Appendix A: CMMC Level 2 Control Mapping Matrix", level=1)
    doc.add_paragraph(
        f"The following matrix maps each {profile['company_short_name']} "
        "policy and procedure to the CMMC Level 2 practices and NIST SP "
        "800-171 Rev. 2 requirements they implement. The standalone Word and "
        "Excel versions of this matrix are part of the CMMC documentation "
        "binder.")
    table = br.styled_table(doc, MATRIX_HEADERS)
    for row in _matrix_rows(profile):
        br.add_table_row(table, row, size=7)


def _appendix_b(doc, profile):
    doc.add_heading(
        "Appendix B: CMMC Level 2 Procedures, Forms, Logs, and Checklists",
        level=1)
    doc.add_paragraph(
        f"The following procedure packages support the "
        f"{profile['company_short_name']} CMMC Level 2 program. Each package "
        "contains the domain procedure and its associated forms, logs, and "
        "checklists, which serve as the operating records of the security "
        "program.")
    table = br.styled_table(doc, ["Domain", "Policy", "Procedure",
                                  "Forms, Logs, and Checklists"])
    for code, name in DOMAINS.items():
        pkg = PACKAGES[code]
        br.add_table_row(table, [
            f"{code} - {name}",
            POLICY_NAMES[code],
            pkg["procedure"]["title"],
            "; ".join(d["title"] for d in pkg["documents"]),
        ], size=8)


# ------------------------------------------------------------- Policies

POLICY_BOILERPLATE_SCOPE = (
    "This policy applies to all employees, contractors, consultants, and any "
    "others granted access to {company_legal_name} information systems, "
    "networks, and data, including all systems that process, store, or "
    "transmit Controlled Unclassified Information (CUI)."
)


def build_policy(profile, domain, out_dir):
    """One branded domain policy modeled on the AVASO policy template."""
    name = POLICY_NAMES[domain]
    plist = practices_by_domain(domain)
    doc = br.new_document()

    br.add_cover_page(doc, profile, name,
                      subtitle="CMMC Level 2 | NIST SP 800-171 Rev. 2")
    br.add_disclosure_statement(doc, profile)
    br.add_page_break(doc)
    br.add_document_history_table(doc, profile)
    br.add_page_break(doc)
    br.add_toc(doc)
    br.add_page_break(doc)

    doc.add_heading("1  Purpose", level=1)
    doc.add_paragraph(fill(
        "{company_legal_name} has adopted the " + DOMAINS[domain] +
        " principles established in NIST SP 800-171 Rev. 2 and the "
        "Cybersecurity Maturity Model Certification (CMMC) Level 2. This "
        "policy defines the requirements that protect the confidentiality of "
        "Controlled Unclassified Information (CUI) within the " +
        DOMAINS[domain] + " domain.", profile))

    doc.add_heading("2  Scope", level=1)
    doc.add_paragraph(fill(POLICY_BOILERPLATE_SCOPE, profile))
    doc.add_paragraph(fill(
        "This policy has been developed by {company_legal_name} in support "
        "of our IT management and security program and the company Standard "
        "Practices and Procedures (SPP).", profile))

    doc.add_heading("3  Policy", level=1)
    for i, p in enumerate(plist, start=1):
        doc.add_heading(f"3.{i}  {p['cmmc_id']} - {p['name']}", level=2)
        doc.add_paragraph(fill(
            "{company_short_name} shall " + p["text"][0].lower() + p["text"][1:],
            profile))
        doc.add_paragraph(fill(
            "Implementation of this requirement is described in the " +
            PACKAGES[domain]["procedure"]["title"] +
            " and evidenced by the records listed in the CMMC Level 2 "
            "Control Mapping Matrix.", profile))

    doc.add_heading("4  Roles and Responsibilities", level=1)
    role, freq = DOMAIN_GOVERNANCE[domain]
    for line in [
        "Senior Management Official ({smo_name}, {smo_title}): approves this "
        "policy and provides resources for its implementation.",
        "Facility Security Officer ({fso_name}): integrates this policy with "
        "the company SPP and security education program.",
        "Insider Threat Program Senior Official ({itpso_name}): oversees "
        "monitoring and incident escalation related to this domain.",
        f"Responsible role for day-to-day execution: {role}.",
        "All personnel: comply with this policy and report violations to "
        "{security_contact_email}.",
    ]:
        doc.add_paragraph(fill(line, profile), style="List Bullet")

    doc.add_heading("5  Compliance and Review", level=1)
    doc.add_paragraph(fill(
        "Violations of this policy may result in disciplinary action up to "
        "and including termination. This policy is reviewed at least "
        "annually by the FSO and ITPSO. Operational review frequency for "
        "this domain: " + freq + ".", profile))

    doc.add_heading("6  What This Document Is For", level=1)
    doc.add_paragraph(fill(DOMAIN_GUIDANCE[domain]["what"], profile))
    doc.add_paragraph(fill(
        "During a CMMC Level 2 assessment, this policy demonstrates "
        "management intent for the " + DOMAINS[domain] + " domain. The "
        "matching procedure and its records demonstrate that the policy is "
        "actually operating.", profile))

    doc.add_heading("7  Evidence Records and Assessment Evidence Needed",
                    level=1)
    doc.add_paragraph(
        "Maintain and be ready to show an assessor the following records:")
    for item in DOMAIN_GUIDANCE[domain]["evidence"]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph(
        "Primary evidence artifacts for this domain: " + evidence_for(domain)
        + ".")

    doc.add_heading("8  Common Assessor Questions", level=1)
    for q in DOMAIN_GUIDANCE[domain]["questions"]:
        doc.add_paragraph(q, style="List Bullet")

    doc.add_heading("9  Compliance References", level=1)
    doc.add_paragraph(fill(COMPLIANCE_STATEMENT, profile))
    for ref in COMPLIANCE_REFERENCES:
        doc.add_paragraph(ref, style="List Bullet")
    for ref in [
        "NIST SP 800-171 Rev. 2, requirements " + ", ".join(
            p["nist_id"] for p in plist),
        fill("{company_short_name} Standard Practices and Procedures (SPP)",
             profile),
        PACKAGES[domain]["procedure"]["title"],
    ]:
        doc.add_paragraph(ref, style="List Bullet")

    doc.add_heading("10  Approval", level=1)
    br.add_approval_block(doc, profile)

    br.add_header(doc, profile, name)
    br.add_page_number_footer(
        doc, f"{profile['company_short_name']} | {name} | Version "
             f"{profile['policy_version']} | Company Sensitive")

    path = os.path.join(_ensure(out_dir),
                        f"{_safe(profile['company_short_name'])}_{_safe(name)}.docx")
    doc.save(path)
    return path


def build_all_policies(profile, out_dir):
    return [build_policy(profile, d, out_dir) for d in DOMAINS]


# ----------------------------------------------------------- Procedures

def _doc_shell(profile, title, subtitle):
    doc = br.new_document()
    br.add_cover_page(doc, profile, title, subtitle=subtitle)
    br.add_disclosure_statement(doc, profile)
    br.add_page_break(doc)
    return doc


def _finish(doc, profile, title, out_path):
    br.add_header(doc, profile, title)
    br.add_page_number_footer(
        doc, f"{profile['company_short_name']} | {title} | Company Sensitive")
    doc.save(out_path)
    return out_path


def build_procedure_package(profile, domain, out_dir):
    """Procedure plus forms, logs, and checklists for one domain."""
    pkg = PACKAGES[domain]
    dom_dir = _ensure(os.path.join(out_dir, f"{domain}_{_safe(DOMAINS[domain])}"))
    paths = []

    # Procedure document, modeled on the AVASO Access Management Procedures.
    proc = pkg["procedure"]
    doc = _doc_shell(profile, proc["title"],
                     f"CMMC Level 2 | {DOMAINS[domain]}")
    br.add_document_history_table(doc, profile)
    doc.add_paragraph()
    doc.add_paragraph("Supporting Controls:").runs[0].bold = True
    for cid in proc["controls"]:
        match = next(p for p in practices() if p["nist_id"] == cid)
        doc.add_paragraph(f"{cid}: {match['text']}", style="List Bullet")
    for i, (title, purpose, steps) in enumerate(proc["sections"], start=1):
        doc.add_heading(f"{i}. {title}", level=1)
        p = doc.add_paragraph()
        p.add_run("Purpose: ").bold = True
        p.add_run(fill(purpose, profile))
        doc.add_paragraph("Procedures:").runs[0].bold = True
        for step in steps:
            doc.add_paragraph(fill(step, profile), style="List Bullet")
    n = len(proc["sections"])
    doc.add_heading(f"{n + 1}. What This Document Is For", level=1)
    doc.add_paragraph(fill(DOMAIN_GUIDANCE[domain]["what"], profile))
    doc.add_heading(f"{n + 2}. Evidence Records and Assessment Evidence "
                    "Needed", level=1)
    doc.add_paragraph(
        "Completed copies of the following are the operating records of "
        "this procedure. Retain them for at least 3 years or as required by "
        "contract, and be ready to show an assessor:")
    for item in DOMAIN_GUIDANCE[domain]["evidence"]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading(f"{n + 3}. Common Assessor Questions", level=1)
    for q in DOMAIN_GUIDANCE[domain]["questions"]:
        doc.add_paragraph(q, style="List Bullet")
    doc.add_heading(f"{n + 4}. Review Frequency", level=1)
    role, freq = DOMAIN_GOVERNANCE[domain]
    doc.add_paragraph(fill(
        "Responsible role: " + role + ". Operational review frequency: "
        + freq + ". This procedure is reviewed at least annually by the FSO "
        "and ITPSO.", profile))
    doc.add_heading(f"{n + 5}. Compliance References", level=1)
    doc.add_paragraph(fill(COMPLIANCE_STATEMENT, profile))
    for ref in COMPLIANCE_REFERENCES:
        doc.add_paragraph(ref, style="List Bullet")
    doc.add_heading(f"{n + 6}. Approval", level=1)
    br.add_approval_block(doc, profile)
    paths.append(_finish(doc, profile, proc["title"], os.path.join(
        dom_dir, f"{_safe(profile['company_short_name'])}_{_safe(proc['title'])}.docx")))

    for spec in pkg["documents"]:
        paths.append(_build_artifact(profile, domain, spec, dom_dir))
    return paths


def _build_artifact(profile, domain, spec, dom_dir):
    title = spec["title"]
    doc = _doc_shell(profile, title, f"CMMC Level 2 | {DOMAINS[domain]}")

    if spec["kind"] == "form":
        doc.add_paragraph(fill(
            "Complete all fields. Submit the signed form to the FSO or IT "
            "security lead as directed by the " +
            PACKAGES[domain]["procedure"]["title"] + ".", profile))
        table = br.styled_table(doc, ["Field", "Entry"])
        for field in spec["fields"]:
            br.add_table_row(table, [fill(field, profile), ""], size=10)
    elif spec["kind"] == "log":
        doc.add_paragraph(fill(
            "Maintain this log as the record of activity required by the " +
            PACKAGES[domain]["procedure"]["title"] +
            ". Retain completed logs for at least 3 years or as required by "
            "contract.", profile))
        table = br.styled_table(doc, spec["columns"])
        for _ in range(12):
            br.add_table_row(table, [""] * len(spec["columns"]), size=9)
    elif spec["kind"] == "checklist":
        doc.add_paragraph(fill(
            "Check each item as it is verified. File the completed checklist "
            "with the records for the " +
            PACKAGES[domain]["procedure"]["title"] + ".", profile))
        for item in spec["items"]:
            br.add_checkbox_item(doc, fill(item, profile))
        doc.add_paragraph()
        for line in ["Completed by: ______________________  Date: ____________",
                     "Reviewed by:  ______________________  Date: ____________"]:
            doc.add_paragraph(line)

    return _finish(doc, profile, title, os.path.join(
        dom_dir, f"{_safe(profile['company_short_name'])}_{_safe(title)}.docx"))


def build_all_procedures(profile, out_dir):
    paths = []
    for domain in DOMAINS:
        paths.extend(build_procedure_package(profile, domain, out_dir))
    return paths


# --------------------------------------------------------------- Matrix

def build_matrix_docx(profile, out_dir):
    doc = br.new_document()
    title = "CMMC Level 2 Control Mapping Matrix"
    br.add_cover_page(doc, profile, title,
                      subtitle="NIST SP 800-171 Rev. 2 | 110 Requirements")
    br.add_disclosure_statement(doc, profile)
    br.add_page_break(doc)

    # Landscape section for the wide table.
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches
    section = doc.add_section()
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11), Inches(8.5)

    doc.add_heading(title, level=1)
    table = br.styled_table(doc, MATRIX_HEADERS)
    for row in _matrix_rows(profile):
        br.add_table_row(table, row, size=7)

    doc.add_heading("Compliance References", level=1)
    doc.add_paragraph(fill(COMPLIANCE_STATEMENT, profile))
    for ref in COMPLIANCE_REFERENCES:
        doc.add_paragraph(ref, style="List Bullet")

    br.add_header(doc, profile, title)
    br.add_page_number_footer(
        doc, f"{profile['company_short_name']} | {title} | Company Sensitive")
    path = os.path.join(_ensure(out_dir),
                        f"{_safe(profile['company_short_name'])}_Control_Mapping_Matrix.docx")
    doc.save(path)
    return path


def build_matrix_xlsx(profile, out_dir):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Control Mapping"
    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append([f"{profile['company_legal_name']} | CMMC Level 2 Control "
               f"Mapping Matrix | CAGE {profile['cage_code']} | "
               f"UEI {profile['uei']}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1,
                   end_column=len(MATRIX_HEADERS))
    ws.cell(1, 1).font = Font(bold=True, size=12)

    ws.append(MATRIX_HEADERS)
    for col in range(1, len(MATRIX_HEADERS) + 1):
        c = ws.cell(2, col)
        c.fill, c.font, c.border = header_fill, header_font, border
        c.alignment = Alignment(wrap_text=True, vertical="center")

    for row in _matrix_rows(profile):
        ws.append(row)
    widths = [22, 16, 14, 30, 30, 32, 45, 18, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(3, ws.max_row + 1):
        for c in range(1, len(MATRIX_HEADERS) + 1):
            cell = ws.cell(r, c)
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(size=9)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(MATRIX_HEADERS))}{ws.max_row}"

    refs = wb.create_sheet("Compliance References")
    refs.append(["Compliance References"])
    refs.cell(1, 1).font = Font(bold=True, size=12)
    refs.append([fill(COMPLIANCE_STATEMENT, profile)])
    for ref in COMPLIANCE_REFERENCES:
        refs.append([ref])
    refs.column_dimensions["A"].width = 120
    for r in range(1, refs.max_row + 1):
        refs.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")

    path = os.path.join(_ensure(out_dir),
                        f"{_safe(profile['company_short_name'])}_Control_Mapping_Matrix.xlsx")
    wb.save(path)
    return path


# --------------------------------------------------------------- Binder

def build_evidence_checklist(profile, out_dir):
    """Assessment evidence checklist across all 14 domains."""
    title = "CMMC Level 2 Evidence Checklist"
    doc = _doc_shell(profile, title,
                     "Assessment Evidence by Domain | NIST SP 800-171 Rev. 2")
    doc.add_paragraph(fill(
        "Use this checklist to confirm that the records an assessor will "
        "ask for actually exist and are current. Check each item, note "
        "where the evidence is stored, and re-run this checklist before "
        "every self-assessment and assessment.", profile))
    for code, name in DOMAINS.items():
        g = DOMAIN_GUIDANCE[code]
        role, freq = DOMAIN_GOVERNANCE[code]
        doc.add_heading(f"{code} - {name}", level=1)
        p = doc.add_paragraph()
        p.add_run("What this domain is for: ").bold = True
        p.add_run(fill(g["what"], profile))
        for item in g["evidence"]:
            br.add_checkbox_item(doc, item)
        p = doc.add_paragraph()
        p.add_run("Responsible role / review frequency: ").bold = True
        p.add_run(f"{role} / {freq}")
        p = doc.add_paragraph()
        p.add_run("Common assessor questions: ").bold = True
        p.add_run(" ".join(g["questions"]))
    doc.add_heading("Compliance References", level=1)
    doc.add_paragraph(fill(COMPLIANCE_STATEMENT, profile))
    for ref in COMPLIANCE_REFERENCES:
        doc.add_paragraph(ref, style="List Bullet")
    doc.add_heading("Approval", level=1)
    br.add_approval_block(doc, profile)
    return _finish(doc, profile, title, os.path.join(
        _ensure(out_dir),
        f"{_safe(profile['company_short_name'])}_Evidence_Checklist.docx"))


def build_binder(profile, out_dir, zip_output=True):
    """Full branded binder: SPP, policies, procedures, matrix, evidence."""
    paths = []
    paths.append(build_spp(profile, os.path.join(out_dir, "01_SPP")))
    paths.extend(build_all_policies(profile, os.path.join(out_dir, "02_Policies")))
    paths.extend(build_all_procedures(profile, os.path.join(out_dir, "03_Procedures")))
    paths.append(build_matrix_docx(profile, os.path.join(out_dir, "04_Control_Matrix")))
    paths.append(build_matrix_xlsx(profile, os.path.join(out_dir, "04_Control_Matrix")))
    paths.append(build_evidence_checklist(profile, os.path.join(out_dir, "05_Evidence")))

    if zip_output:
        import zipfile
        zip_path = os.path.join(
            out_dir, f"{_safe(profile['company_short_name'])}_CMMC_L2_Binder.zip")
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
            for p in paths:
                zf.write(p, os.path.relpath(p, out_dir))
        paths.append(zip_path)
    return paths
